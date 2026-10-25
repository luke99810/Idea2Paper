"""Render the single-source manuscript.md into LaTeX and Word versions.
Mechanical conversion only; never invents content. LaTeX PDF requires xelatex;
Word requires python-docx. Missing tools are reported, not faked.
Usage:
  python scripts/render.py latex <project>   -> draft/manuscript.tex (+ PDF if xelatex found)
  python scripts/render.py docx  <project>   -> draft/manuscript.docx
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import svg2png

NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def find_project(root):
    path = Path(root).resolve()
    md = path / 'draft' / 'manuscript.md'
    if not md.exists():
        raise ValueError('Missing draft/manuscript.md in project: ' + str(path))
    return path, md


# ---------- shared markdown block parsing ----------

def split_blocks(text):
    """Parse the manuscript subset, binding captions to their owning objects."""
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        heading = re.match(r'^(#{1,6})\s+(.*)$', line)
        image = re.fullmatch(r'!\[(.*?)\]\((.+?)\)', line)
        if heading:
            blocks.append(('h', (len(heading[1]), heading[2])))
        elif image:
            blocks.append(('image', (image[2], image[1])))
        elif line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            caption = ''
            if blocks and blocks[-1][0] == 'p' and re.fullmatch(r'\*\*表\s*\d+.*\*\*', blocks[-1][1]):
                caption = blocks.pop()[1][2:-2]
            blocks.append(('table', (caption, rows)))
        elif re.fullmatch(r'\*\*图\s*\d+.*\*\*', line) and blocks and blocks[-1][0] == 'image':
            blocks[-1] = ('image', (blocks[-1][1][0], line[2:-2]))
        elif line.startswith('$$'):
            math = line[2:]
            while not math.rstrip().endswith('$$') and i < len(lines):
                math += '\n' + lines[i]
                i += 1
            if not math.rstrip().endswith('$$'):
                raise ValueError('Unclosed display math')
            blocks.append(('math', math.rstrip()[:-2].strip()))
        elif re.match(r'^([-*+]|\d+[.、])\s+', line):
            item = re.sub(r'^([-*+]|\d+[.、])\s+', '', line)
            if blocks and blocks[-1][0] == 'list':
                blocks[-1][1].append(item)
            else:
                blocks.append(('list', [item]))
        elif line == '---':
            blocks.append(('hr', None))
        else:
            blocks.append(('p', line))
    return blocks

def _is_sep_row(cells):
    """A markdown table separator row: every cell matches ^:?-+:?$."""
    return bool(cells) and all(re.fullmatch(r':?-+:?', c) for c in cells)


def _table_rows(data):
    rows = [r.strip().strip('|').split('|') for r in data]
    cells = [[c.strip() for c in row] for row in rows]
    cells = [r for r in cells if not _is_sep_row(r)]
    return cells


def inline_latex(text):
    tokens = re.split(r'(\$[^$\n]+\$|\\\(.+?\\\)|\\\[.+?\\\]|\*\*.+?\*\*|\*[^*]+\*)', text)
    escapes = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%',
               '$': r'\$', '#': r'\#', '_': r'\_', '{': r'\{',
               '}': r'\}', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    out = []
    for part in tokens:
        if part.startswith(('$', r'\(', r'\[')) and len(part) > 1:
            out.append(part)
        elif part.startswith('**') and part.endswith('**'):
            out.append(r'\textbf{' + inline_latex(part[2:-2]) + '}')
        elif part.startswith('*') and part.endswith('*'):
            out.append(r'\textit{' + inline_latex(part[1:-1]) + '}')
        else:
            out.append(''.join(escapes.get(c, c) for c in part))
    return ''.join(out)


# ---------- LaTeX ----------

def resolve_citations(text, records):
    """Resolve author/year labels from supplied metadata, never guessed authors."""
    by_id = {record['id']: record for record in records}
    def replace(match):
        labels = []
        for key in re.findall(r'@([^;\s\]]+)', match[0]):
            if key not in by_id:
                raise ValueError('Unknown citation id: ' + key)
            record = by_id[key]
            citation = record.get('citation', '')
            year = str(record.get('year', ''))
            authors = record.get('author', record.get('authors', []))
            if isinstance(authors, str):
                authors = [authors]
            names = [a.get('family', a.get('literal', '')) if isinstance(a, dict) else a for a in authors]
            if not year:
                found = re.search(r'\((\d{4}[a-z]?)\)', citation)
                year = found[1] if found else ''
            if not names and year:
                prefix = citation.split('(' + year + ')')[0].strip().rstrip('.')
                names = re.findall(r'(?:^|&\s*|and\s+)([^,]+),', prefix)
            names = [name.strip() for name in names if name.strip()]
            if not names or not year:
                raise ValueError('Citation lacks author/year metadata: ' + key)
            label = ' & '.join(names[:2]) if len(names) <= 2 else names[0] + ' et al.'
            labels.append(label + ', ' + year)
        return '(' + '; '.join(labels) + ')'
    return re.sub(r'\[@[^\]]+\]', replace, text)


def manuscript_text(root, text):
    path = Path(root) / 'bibliography.json' if root is not None else None
    records = json.loads(path.read_text(encoding='utf-8')) if path and path.exists() else []
    text = re.sub(r"(?m)^> ?", "", text)
    config_path = Path(root) / 'render-config.json' if root is not None else None
    config = json.loads(config_path.read_text(encoding='utf-8')) if config_path and config_path.exists() else {}
    if config.get('citation_style') != 'numeric':
        return resolve_citations(text, records)
    by_id = {record['id']: record for record in records}
    order = []
    def numbered(match):
        numbers = []
        for key in re.findall(r'@([^;\s\]]+)', match[0]):
            if key not in by_id:
                raise ValueError('Unknown citation id: ' + key)
            if key not in order:
                order.append(key)
            numbers.append(str(order.index(key) + 1))
        return '[' + ','.join(numbers) + ']'
    text = re.sub(r'\[@[^\]]+\]', numbered, text)
    entries = ['[' + str(i + 1) + '] ' + by_id[key]['citation'] for i, key in enumerate(order)]
    pattern = r'(?ms)(^## (?:参考文献|References)\s*\n).*?(?=^## |\Z)'
    return re.sub(pattern, lambda match: match[1] + '\n' + '\n\n'.join(entries) + '\n\n', text)


def prepare_image(root, path):
    """Require a real local image; refresh SVG rasterizations on every render."""
    image = (Path(root) / path).resolve()
    if not image.is_file():
        raise FileNotFoundError('Missing manuscript image: ' + str(image))
    if image.suffix.lower() == '.svg':
        png = image.with_suffix('.png')
        png.write_bytes(svg2png.svg2png(image.read_text(encoding='utf-8'), scale=4))
        image = png
    return image


def to_latex(md_text, template=None, root=None):
    if template is None:
        template = (Path(__file__).resolve().parents[1] /
                    'templates/latex/manuscript.tex').read_text(encoding='utf-8')
    if '@@TITLE@@' not in template or '@@BODY@@' not in template:
        raise ValueError('LaTeX template requires @@TITLE@@ and @@BODY@@ placeholders')
    md_text = manuscript_text(root, md_text)
    title_text = ''
    out = []
    in_abstract = False
    for kind, data in split_blocks(md_text):
        if kind == 'h':
            level, title = data
            if level == 1 and not title_text:
                title_text = inline_latex(title)
                continue
            if title.strip() in ('摘要', 'Abstract'):
                out.append('\\begin{abstract}')
                in_abstract = True
            elif title.strip() in ('关键词', 'Keywords'):
                continue
            else:
                if in_abstract:
                    out.append('\\end{abstract}')
                    in_abstract = False
                cmd = {1: 'section', 2: 'section', 3: 'subsection', 4: 'subsubsection'}.get(level, 'paragraph')
                out.append(f'\\{cmd}{{{inline_latex(title)}}}')
        elif kind == 'p':
            text = data.strip()
            if not text:
                continue
            if in_abstract:
                out.append(inline_latex(text))
            else:
                out.append(inline_latex(text) + '\n')
        elif kind == 'table':
            cells = _table_rows(data[1])
            cols = max(len(r) for r in cells)
            header = cells[0]
            body = cells[1:]
            out.append('\\begin{table}[htbp]\n\\centering')
            if data[0]:
                out.append(r'\caption*{' + inline_latex(data[0]) + '}')
            out.append('\\begin{tabular}{' + 'c' * cols + '}')
            out.append('\\toprule\n' + ' & '.join(inline_latex(c) for c in header) + r' \\')
            out.append('\\midrule')
            for row in body:
                out.append(' & '.join(inline_latex(c) for c in row[:cols]) + r' \\')
            out.append('\\bottomrule\n\\end{tabular}\n\\end{table}')
        elif kind == 'list':
            out.append('\\begin{itemize}')
            for item in data:
                out.append('\\item ' + inline_latex(item))
            out.append('\\end{itemize}')
        elif kind == 'image':
            image = prepare_image(Path(root) if root is not None else Path.cwd(), data[0])
            path = image.relative_to(Path(root).resolve()).as_posix() if root is not None else image.as_posix()
            out.append('\\begin{figure}[htbp]\n\\centering\n' +
                       r'\includegraphics[width=\linewidth,height=0.8\textheight,keepaspectratio]{\detokenize{' + path + '}}')
            if data[1]:
                out.append(r'\caption*{' + inline_latex(data[1]) + '}')
            out.append(r'\end{figure}')
        elif kind == 'math':
            out.append('\\begin{equation*}\n' + data.strip('$') + '\n\\end{equation*}')
    if in_abstract:
        out.append('\\end{abstract}')
    return template.replace('@@TITLE@@', title_text).replace('@@BODY@@', '\n'.join(out))


def render_latex(project):
    root, md = find_project(project)
    tex = md.with_suffix('.tex')
    tex.write_text(to_latex(md.read_text(encoding='utf-8'), root=root), encoding='utf-8')
    target = tex.with_suffix('.pdf')
    target.unlink(missing_ok=True)
    result = {'tex': tex.relative_to(root).as_posix(), 'pdf': None}
    xelatex = shutil.which('xelatex')
    if not xelatex:
        result['note'] = 'xelatex not found; LaTeX source kept, PDF not compiled'
        return result
    with tempfile.TemporaryDirectory(prefix='idea2paper-tex-') as directory:
        build = Path(directory)
        run = subprocess.run([xelatex, '-interaction=nonstopmode', '-halt-on-error',
                              '-output-directory', str(build), str(tex)],
                             cwd=str(root), capture_output=True, text=True,
                             encoding='utf-8', errors='replace', check=False)
        pdf = build / (tex.stem + '.pdf')
        if run.returncode != 0 or not pdf.is_file() or not pdf.read_bytes().startswith(b'%PDF-'):
            result.update(status='fail', note='xelatex failed or produced no new PDF',
                          returncode=run.returncode, log=run.stdout + run.stderr)
            return result
        target.write_bytes(pdf.read_bytes())
    result['pdf'] = target.relative_to(root).as_posix()
    return result


# ---------- Word (python-docx) ----------

def add_inline_docx(paragraph, text, font_cn='宋体', font_en='Times New Roman', size=12, bold=False):
    from docx.shared import Pt
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        b = bold
        if part.startswith('**') and part.endswith('**'):
            part = part[2:-2]
            b = True
        run = paragraph.add_run(part)
        run.font.name = font_en
        run._element.rPr.rFonts.set(NS + 'eastAsia', font_cn)
        run.font.size = Pt(size)
        run.bold = b


def three_line_table(table):
    """Use only outer top/bottom and header-bottom rules, with repeated headers."""
    from docx.oxml import OxmlElement
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        element.set(NS + 'val', 'single' if edge in ('top', 'bottom') else 'nil')
        element.set(NS + 'sz', '8')
        borders.append(element)
    table._tbl.tblPr.append(borders)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in table.rows[0].cells:
        cell_borders = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(NS + 'val', 'single')
        bottom.set(NS + 'sz', '6')
        cell_borders.append(bottom)
        cell._tc.get_or_add_tcPr().append(cell_borders)


def to_docx(root, md_text, md_path):
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement

    md_text = manuscript_text(root, md_text)
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(NS + 'eastAsia', '宋体')
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.first_line_indent = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    for s in doc.sections:
        s.page_width = Cm(21)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.54)
        s.right_margin = Cm(2.54)

    title_seen = False
    in_abstract = False
    for kind, data in split_blocks(md_text):
        if kind == 'h':
            level, title = data
            in_abstract = title.strip() in ('摘要', 'Abstract')
            is_title = level == 1 and not title_seen
            title_seen = title_seen or is_title
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            if is_title:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(title)
            r.bold = True
            r.font.name = 'Times New Roman'
            r._element.rPr.rFonts.set(NS + 'eastAsia', '黑体')
            r.font.size = Pt(22 if is_title else 12)
        elif kind == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Pt(0) if in_abstract or not title_seen or not data.strip() else Cm(0.84)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if re.match(r'^\[\d+\] ', data.strip()):
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.first_line_indent = Cm(-0.7)
            elif data.strip().startswith(('注：', '**关键词', '教学演示声明')):
                p.paragraph_format.first_line_indent = Pt(0)
            add_inline_docx(p, data.strip())
        elif kind == 'table':
            cells = _table_rows(data[1])
            rows_n = max(len(r) for r in cells)
            if data[0]:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.keep_with_next = True
                add_inline_docx(cap, data[0], size=10.5)
            table = doc.add_table(rows=len(cells), cols=rows_n)
            table.autofit = False
            usable_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
            weights = [1.25, 0.85, 1.8, 1.15, 3.05, 1.15, 1.3, 2.8] if rows_n == 8 else [1] * rows_n
            widths = [int(usable_width * weight / sum(weights)) for weight in weights]
            for column, width in zip(table.columns, widths):
                column.width = width
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = width
                    cell.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
                    cell.paragraphs[0].paragraph_format.space_before = Pt(0)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            three_line_table(table)
            for i, row in enumerate(cells):
                for j in range(rows_n):
                    cell = table.cell(i, j)
                    if i > 0:
                        cell._tc.get_or_add_tcPr().append(OxmlElement('w:noWrap'))
                    cell.paragraphs[0].paragraph_format.keep_with_next = i < len(cells) - 1 and len(cells) <= 8
                    cell.paragraphs[0].paragraph_format.keep_together = True
                    add_inline_docx(cell.paragraphs[0], row[j] if j < len(row) else '', size=8.5)
                    if i == 0:
                        cell.paragraphs[0].runs and setattr(cell.paragraphs[0].runs[0], 'bold', True)
        elif kind == 'list':
            for item in data:
                p = doc.add_paragraph(style='List Bullet')
                add_inline_docx(p, item)
        elif kind == 'image':
            img = prepare_image(root, data[0])
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = bool(data[1])
            p.paragraph_format.keep_together = True
            section = doc.sections[-1]
            width = section.page_width - section.left_margin - section.right_margin
            height = section.page_height - section.top_margin - section.bottom_margin - Cm(2)
            picture = p.add_run().add_picture(str(img), width=width)
            if picture.height > height:
                picture.width = int(picture.width * height / picture.height)
                picture.height = height
            if data[1]:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.keep_together = True
                add_inline_docx(cap, data[1], size=10.5)
        elif kind == 'math':
            p = doc.add_paragraph()
            add_inline_docx(p, data.strip('$'))
    out = root / 'draft' / 'manuscript.docx'
    doc.save(str(out))
    return out.name


def render_docx(project):
    root, md = find_project(project)
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return {'status': 'fail', 'error': 'python-docx required: pip install python-docx'}
    name = to_docx(root, md.read_text(encoding='utf-8'), md)
    return {'docx': name}


def main():
    args = sys.argv[1:]
    if len(args) != 2 or args[0] not in ('latex', 'docx'):
        print(__doc__, file=sys.stderr)
        return 2
    try:
        result = (render_latex if args[0] == 'latex' else render_docx)(args[1])
        print(result)
        return 1 if result.get('status') == 'fail' else 0
    except Exception as exc:
        print({'status': 'fail', 'error': str(exc)}, file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
