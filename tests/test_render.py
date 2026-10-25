"""Offline rendering regression tests; all fixtures are synthetic."""
import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import render


class RenderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / 'draft').mkdir()
        self.md = self.root / 'draft/manuscript.md'
        self.md.write_text('# Synthetic title\n\nBody.', encoding='utf-8')

    def test_svg_is_print_resolution_at_manuscript_width(self):
        from PIL import Image
        svg = self.root / 'figure.svg'
        svg.write_text('<svg width="800" height="460"><rect width="800" height="460" fill="white"/></svg>')
        image = render.prepare_image(self.root, 'figure.svg')
        with Image.open(image) as picture:
            self.assertGreaterEqual(picture.width / (15.92 / 2.54), 300)

    def test_abstract_has_no_first_line_indent(self):
        text = '# T\n\n## 摘要\n\nAbstract body.'
        render.to_docx(self.root, text, self.md)
        from docx import Document
        paragraph = Document(self.md.with_suffix('.docx')).paragraphs[2]
        self.assertEqual(paragraph.paragraph_format.first_line_indent.pt, 0)

    def test_blockquote_marker_is_not_literal_prose(self):
        render.to_docx(self.root, '# T\n\n> Teaching notice', self.md)
        from docx import Document
        doc = Document(self.md.with_suffix('.docx'))
        self.assertEqual(doc.paragraphs[1].text, 'Teaching notice')

    def test_table_rows_stay_with_header_in_short_table(self):
        render.to_docx(self.root, '# T\n\n|A|B|\n|---|---|\n|1|2|', self.md)
        from docx import Document
        table = Document(self.md.with_suffix('.docx')).tables[0]
        self.assertTrue(table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)

    def test_inline_escapes_text_without_corrupting_math_or_emphasis(self):
        text = r'**bold & 50%** *italics* a_b # {x} $x_1 + \alpha$ \(y^2\)'
        self.assertEqual(render.inline_latex(text),
                         r'\textbf{bold \& 50\%} \textit{italics} a\_b \# \{x\} $x_1 + \alpha$ \(y^2\)')

    def test_latex_uses_complete_replaceable_template_and_title(self):
        template = ('\\documentclass{ctexart}\n% custom skeleton\n'
                    '\\usepackage{amsmath,graphicx,booktabs}\n'
                    '\\begin{document}\n\\begin{center}@@TITLE@@\\end{center}\n'
                    '@@BODY@@\n\\end{document}')
        tex = render.to_latex('# Title & test\n\n## Introduction\n\nBody.', template=template)
        self.assertIn('% custom skeleton', tex)
        self.assertIn(r'\begin{center}Title \& test\end{center}', tex)
        self.assertIn(r'\section{Introduction}', tex)
        self.assertNotIn(r'\section{Title', tex)
        self.assertNotIn('@@', tex)
        self.assertTrue(tex.endswith(r'\end{document}'))
        with self.assertRaisesRegex(ValueError, 'placeholder'):
            render.to_latex('# Title', template='invalid')

    def test_blocks_bind_unique_figure_and_table_captions(self):
        blocks = render.split_blocks('# T\nText\n![](figures/a.svg)\n\n**图 2：唯一图注**\n\n'
                                     '**表 1：表题**\n\n|A|B|\n|---|---|\n|1|2|\n\n$$x_1=1$$\nAfter')
        self.assertEqual(blocks[2], ('image', ('figures/a.svg', '图 2：唯一图注')))
        self.assertEqual(blocks[3][0], 'table')
        self.assertEqual(blocks[3][1][0], '表 1：表题')
        self.assertEqual(blocks[4], ('math', 'x_1=1'))
        self.assertEqual(blocks[5], ('p', 'After'))
        self.assertEqual(render._table_rows(blocks[3][1][1]), [['A', 'B'], ['1', '2']])
        self.assertTrue(render._is_sep_row([':---', '---:']))

    def test_word_title_and_page_geometry(self):
        from docx import Document
        from docx.shared import Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        render.to_docx(self.root, '# 独立题名\n\n## 引言\n\n正文 **加粗**', self.md)
        doc = Document(self.md.with_suffix('.docx'))
        title = doc.paragraphs[0]
        self.assertTrue(title.paragraph_format.keep_with_next)
        self.assertEqual(title.paragraph_format.first_line_indent.pt, 0)
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(title.runs[0].font.size.pt, 22)
        self.assertEqual(title.runs[0]._element.rPr.rFonts.get(render.NS + 'eastAsia'), '黑体')
        section = doc.sections[0]
        for attr, value in [('page_width', 21), ('page_height', 29.7),
                            ('left_margin', 2.54), ('right_margin', 2.54),
                            ('top_margin', 2.54), ('bottom_margin', 2.54)]:
            self.assertAlmostEqual(getattr(section, attr).cm, value, places=2)
        self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
        self.assertEqual(doc.styles['Normal'].paragraph_format.line_spacing, 1.5)
        headings = [para for para in doc.paragraphs if para.text == '引言']
        self.assertEqual(len(headings), 1)
        self.assertEqual(headings[0].paragraph_format.first_line_indent.pt, 0)
        body = doc.paragraphs[2]
        self.assertAlmostEqual(body.paragraph_format.first_line_indent.cm, 0.84, places=2)
        self.assertEqual(body.paragraph_format.space_before.pt, 0)
        self.assertEqual(body.paragraph_format.space_after.pt, 0)
        self.assertTrue(doc.paragraphs[2].runs[-1].bold)

    def test_figures_convert_svg_and_keep_unique_caption_with_image(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        figures = self.root / 'figures'
        figures.mkdir()
        (figures / 'wide.svg').write_text('<svg width="1200" height="200"><rect width="1200" height="200" fill="white"/></svg>', encoding='utf-8')
        text = '# T\n\n![](figures/wide.svg)\n\n**图 1：唯一图题**'
        render.to_docx(self.root, text, self.md)
        doc = Document(self.md.with_suffix('.docx'))
        self.assertEqual([p.text for p in doc.paragraphs].count('图 1：唯一图题'), 1)
        picture = doc.paragraphs[1]
        self.assertEqual(picture.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(picture.paragraph_format.keep_with_next)
        self.assertLessEqual(doc.inline_shapes[0].width, doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin)
        self.assertTrue((figures / 'wide.png').is_file())
        template = r'\documentclass{ctexart}\begin{document}@@TITLE@@@@BODY@@\end{document}'
        tex = render.to_latex(text, template=template, root=self.root)
        self.assertEqual(tex.count('唯一图题'), 1)
        self.assertIn('figures/wide.png', tex)
        self.assertNotIn(r'\caption{}', tex)
        self.assertIn(r'\linewidth', tex)
        for target in ['docx', 'latex']:
            with self.subTest(target=target), self.assertRaisesRegex(FileNotFoundError, 'missing'):
                text = '# T\n\n![](figures/missing.svg)'
                if target == 'docx':
                    render.to_docx(self.root, text, self.md)
                else:
                    render.to_latex(text, template=template, root=self.root)

    def test_three_line_tables_have_caption_above_and_no_grid(self):
        from docx import Document
        text = '# T\n\n**表 1：真实表题**\n\n|A|B|\n|---|---|\n|1|2|'
        render.to_docx(self.root, text, self.md)
        doc = Document(self.md.with_suffix('.docx'))
        self.assertEqual(doc.paragraphs[1].text, '表 1：真实表题')
        self.assertTrue(doc.paragraphs[1].paragraph_format.keep_with_next)
        table = doc.tables[0]
        self.assertNotEqual(table.style.name, 'Table Grid')
        borders = table._tbl.tblPr.find(render.NS + 'tblBorders')
        self.assertEqual(borders.find(render.NS + 'top').get(render.NS + 'val'), 'single')
        self.assertEqual(borders.find(render.NS + 'bottom').get(render.NS + 'val'), 'single')
        self.assertEqual(borders.find(render.NS + 'insideV').get(render.NS + 'val'), 'nil')
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(render.NS + 'tblHeader'))
        self.assertIsNotNone(table.cell(0, 0)._tc.tcPr.find(render.NS + 'tcBorders'))
        self.assertFalse(table.autofit)
        self.assertTrue(all(cell.paragraphs[0].paragraph_format.first_line_indent.pt == 0
                            for row in table.rows for cell in row.cells))
        self.assertTrue(all(run.font.size.pt <= 9 for row in table.rows for cell in row.cells
                            for run in cell.paragraphs[0].runs))
        self.assertTrue(all(cell._tc.tcPr.find(render.NS + 'noWrap') is not None
                            for row in table.rows[1:] for cell in row.cells))
        template = '@@TITLE@@\n@@BODY@@'
        tex = render.to_latex(text, template=template)
        self.assertLess(tex.index('真实表题'), tex.index(r'\begin{tabular}'))
        self.assertIn(r'\toprule', tex)
        self.assertIn(r'\midrule', tex)
        self.assertIn(r'\bottomrule', tex)
        self.assertNotIn('占位', tex)

    def test_numbered_references_follow_first_citation_and_have_no_bullets(self):
        from docx import Document
        records = [{'id': 'b', 'citation': 'B reference.'}, {'id': 'a', 'citation': 'A reference.'}]
        (self.root / 'bibliography.json').write_text(json.dumps(records))
        (self.root / 'render-config.json').write_text(json.dumps({'citation_style': 'numeric'}))
        text = '# T\n\n## 引言\n\nFirst [@a], then [@b; @a].\n\n## 参考文献\n\n- Old entry'
        render.to_docx(self.root, text, self.md)
        doc = Document(self.md.with_suffix('.docx'))
        self.assertIn('First [1], then [2,1].', [p.text for p in doc.paragraphs])
        self.assertIn('[1] A reference.', [p.text for p in doc.paragraphs])
        self.assertIn('[2] B reference.', [p.text for p in doc.paragraphs])
        self.assertFalse(any(p.style.name == 'List Bullet' for p in doc.paragraphs))

    def test_citations_resolve_only_from_real_bibliography_metadata(self):
        from docx import Document
        records = [{'id': 'known', 'citation': 'Ropke, S., & Pisinger, D. (2006). Real title.'},
                   {'id': 'structured', 'author': [{'family': 'Li', 'given': 'Ming'}], 'year': 2024}]
        (self.root / 'bibliography.json').write_text(json.dumps(records), encoding='utf-8')
        text = '# T\n\nEvidence [@known; @structured].'
        tex = render.to_latex(text, template='@@TITLE@@@@BODY@@', root=self.root)
        self.assertIn('Ropke', tex)
        self.assertIn('2006', tex)
        self.assertIn('Li', tex)
        self.assertNotIn('[@', tex)
        render.to_docx(self.root, text, self.md)
        self.assertNotIn('[@', '\n'.join(p.text for p in Document(self.md.with_suffix('.docx')).paragraphs))
        with self.assertRaisesRegex(ValueError, 'unknown'):
            render.to_latex('# T\n\n[@unknown]', template='@@TITLE@@@@BODY@@', root=self.root)
        with self.assertRaisesRegex(ValueError, 'metadata'):
            render.resolve_citations('[@empty]', [{'id': 'empty'}])
        self.assertEqual(render.resolve_citations('No citation.', []), 'No citation.')

    def test_failed_compile_never_reuses_stale_pdf(self):
        (self.root / 'build').mkdir()
        (self.root / 'build/manuscript.pdf').write_bytes(b'OLD BUILD')
        self.md.with_suffix('.pdf').write_bytes(b'OLD DRAFT')
        template = r'\documentclass{ctexart}\begin{document}@@TITLE@@@@BODY@@\end{document}'
        real_to_latex = render.to_latex
        with patch.object(render.shutil, 'which', return_value='fake-xelatex'), \
             patch.object(render, 'to_latex', side_effect=lambda text, **kw: real_to_latex(text, template=template, **kw)), \
             patch.object(render.subprocess, 'run', return_value=subprocess_result(1)):
            result = render.render_latex(self.root)
        self.assertIsNone(result['pdf'])
        self.assertEqual(result['status'], 'fail')
        self.assertFalse(self.md.with_suffix('.pdf').exists())
        self.assertTrue(self.md.with_suffix('.tex').read_text(encoding='utf-8').startswith(r'\documentclass'))

    def test_default_template_produces_complete_source_without_latex_tool(self):
        with patch.object(render.shutil, 'which', return_value=None):
            result = render.render_latex(self.root)
        tex = self.md.with_suffix('.tex').read_text(encoding='utf-8')
        self.assertIn(r'\documentclass[12pt,a4paper]{ctexart}', tex)
        self.assertIn(r'\begin{document}', tex)
        self.assertIn(r'\end{document}', tex)
        self.assertIn('Synthetic title', tex)
        self.assertNotIn('@@', tex)
        self.assertIsNone(result['pdf'])
        self.assertFalse(self.md.with_suffix('.pdf').exists())

    def test_render_entrypoints_and_compiler_output_checks(self):
        from docx import Document
        self.assertEqual(render.find_project(self.root), (self.root.resolve(), self.md.resolve()))
        with self.assertRaises(ValueError):
            render.find_project(self.root / 'missing')
        result = render.render_docx(self.root)
        self.assertEqual(result['docx'], 'manuscript.docx')
        self.assertEqual(Document(self.md.with_suffix('.docx')).paragraphs[0].text, 'Synthetic title')
        template = r'\documentclass{ctexart}\begin{document}@@TITLE@@@@BODY@@\end{document}'
        convert = render.to_latex
        with patch.object(render, 'to_latex', side_effect=lambda text, **kw: convert(text, template=template, **kw)):
            with patch.object(render.shutil, 'which', return_value=None):
                result = render.render_latex(self.root)
                self.assertIsNone(result['pdf'])
                self.assertIn('not found', result['note'])
            with patch.object(render.shutil, 'which', return_value='fake-xelatex'):
                with patch.object(render.subprocess, 'run', return_value=subprocess_result(0)):
                    self.assertEqual(render.render_latex(self.root)['status'], 'fail')
                def compile_fixture(command, **kwargs):
                    self.assertEqual(Path(kwargs['cwd']), self.root)
                    output = Path(command[command.index('-output-directory') + 1]) / 'manuscript.pdf'
                    output.write_bytes(b'%PDF-synthetic-unit-test')
                    return subprocess_result(0)
                with patch.object(render.subprocess, 'run', side_effect=compile_fixture):
                    self.assertEqual(render.render_latex(self.root)['pdf'], 'draft/manuscript.pdf')
                    self.assertEqual(self.md.with_suffix('.pdf').read_bytes(), b'%PDF-synthetic-unit-test')
        with patch.object(sys, 'argv', ['render.py']), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(render.main(), 2)
        with patch.object(sys, 'argv', ['render.py', 'docx', str(self.root)]), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(render.main(), 0)
        with patch.object(sys, 'argv', ['render.py', 'docx', str(self.root / 'missing')]), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(render.main(), 1)

    def test_math_lists_and_empty_alt_have_no_fabricated_caption(self):
        text = '# T\n\n## Abstract\n\nSummary\n\n## Keywords\n\nWords\n\n## Main\n\n- A\n- B\n\n$$\n\\min_{x} x_1\n$$\n\n---'
        tex = render.to_latex(text, template='@@TITLE@@@@BODY@@')
        self.assertIn(r'\begin{abstract}', tex)
        self.assertIn(r'\end{abstract}', tex)
        self.assertIn(r'\item A', tex)
        self.assertIn(r'\min_{x} x_1', tex)
        render.to_docx(self.root, text, self.md)
        self.assertTrue(self.md.with_suffix('.docx').exists())
        with self.assertRaisesRegex(ValueError, 'Unclosed'):
            render.split_blocks('$$\nunclosed')


def subprocess_result(code):
    import subprocess
    return subprocess.CompletedProcess([], code, stdout='', stderr='synthetic compiler failure')


if __name__ == '__main__':
    unittest.main()
